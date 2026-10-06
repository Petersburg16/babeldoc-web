"""会议记录导出：Word、Markdown、纯文本、字幕。

只依赖传进来的 ExportData，不碰数据库，渲染可以放到会话关闭之后做。
说话人占位符 [[S3]] 在这里换成显示名，规则与前端 lib/meeting/format.ts 的 speakerName 一致。
"""

from __future__ import annotations

import io
import re
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import UTC, datetime
from typing import Any, Literal
from urllib.parse import quote

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, RGBColor
from docx.text.paragraph import Paragraph as DocxParagraph

from ..services import LOCAL_TZ
from .format import resolve_speaker, speaker_sort_key, spoken

ExportFormat = Literal["docx", "md", "txt", "srt"]
ExportContent = Literal["minutes", "transcript", "both"]

MEDIA_TYPES: dict[str, str] = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "md": "text/markdown; charset=utf-8",
    "txt": "text/plain; charset=utf-8",
    "srt": "application/x-subrip; charset=utf-8",
}
# Windows 上的老播放器和记事本靠 BOM 认出 UTF-8，不带会把中文显示成乱码
BOM = "﻿"
CJK_FONT = "微软雅黑"
LATIN_FONT = "Calibri"
GRAY = RGBColor(0x80, 0x80, 0x80)

SPEAKER_PLACEHOLDER = re.compile(r"\[\[(S\d+)\]\]")
TIMESTAMP = re.compile(r"\[\d{1,2}:\d{2}(?::\d{2})?\]")


@dataclass
class ExportSegment:
    start_ms: int
    end_ms: int
    speaker: str
    text: str


@dataclass
class ExportData:
    title: str
    created_at: datetime | None = None
    duration_ms: int = 0
    provider_name: str = ""
    speakers: dict[str, Any] = field(default_factory=dict)
    minutes_md: str | None = None
    segments: list[ExportSegment] = field(default_factory=list)


@dataclass
class Turn:
    """同一说话人连续发言合成的一段。"""

    speaker: str
    start_ms: int
    text: str


class NothingToExport(Exception):
    """要导出的内容还不存在，message 直接给用户看。"""


# ---------- 说话人 ----------


def speaker_name(speakers: dict[str, Any], sid: str) -> str:
    final = resolve_speaker(speakers, sid)
    info = speakers.get(final)
    name = str(info.get("name") or "").strip() if isinstance(info, dict) else ""
    return name or f"说话人 {final.removeprefix('S')}"


def fill_speakers(text: str, speakers: dict[str, Any]) -> str:
    return SPEAKER_PLACEHOLDER.sub(lambda m: speaker_name(speakers, m.group(1)), text)


def participants(data: ExportData) -> list[str]:
    """参会人：按首次发言顺序；没有逐字稿时按编号列出未被合并的说话人。"""
    ids = [s.speaker for s in data.segments] or sorted(data.speakers, key=speaker_sort_key)
    seen: dict[str, None] = {}
    for sid in ids:
        seen.setdefault(resolve_speaker(data.speakers, sid), None)
    return [speaker_name(data.speakers, sid) for sid in seen]


# ---------- 时间 ----------


def clock(ms: int) -> str:
    total = max(0, ms) // 1000
    return f"{total // 3600:02d}:{total % 3600 // 60:02d}:{total % 60:02d}"


def srt_clock(ms: int) -> str:
    ms = max(0, ms)
    return f"{clock(ms)},{ms % 1000:03d}"


def local_date(value: datetime | None) -> str:
    if value is None:
        return ""
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value.astimezone(LOCAL_TZ).strftime("%Y-%m-%d")


# ---------- 内容整理 ----------


_XML_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def xml_safe(text: str) -> str:
    """去掉 Word（XML）不允许的控制字符；老数据里可能有从别处粘贴带进来的。"""
    return _XML_ILLEGAL.sub(" ", text)


def _clean(text: str) -> str:
    return re.sub(r"\s*\n\s*", " ", text).strip()


def _join(a: str, b: str) -> str:
    # 中文句子直接相连；两边都是英文数字时补一个空格，免得单词粘在一起
    if a and b and a[-1].isascii() and not a[-1].isspace() and b[0].isascii():
        return f"{a} {b}"
    return a + b


def turns(data: ExportData) -> list[Turn]:
    """按“合并后的”说话人把连续发言合成一段，被合并的人和目标人相邻的话会连在一起。"""
    out: list[Turn] = []
    last: str | None = None
    for seg in data.segments:
        text = _clean(seg.text)
        if not text:
            continue
        sid = resolve_speaker(data.speakers, seg.speaker)
        if out and sid == last:
            out[-1].text = _join(out[-1].text, text)
        else:
            out.append(Turn(sid, seg.start_ms, text))
            last = sid
    return out


def header_lines(data: ExportData) -> list[tuple[str, str]]:
    rows = [
        ("日期", local_date(data.created_at)),
        ("时长", spoken(data.duration_ms) if data.duration_ms > 0 else ""),
        ("参会人", "、".join(participants(data))),
        ("识别服务", data.provider_name.strip()),
    ]
    return [(label, value) for label, value in rows if value]


def sections(data: ExportData, content: ExportContent) -> tuple[str | None, list[Turn] | None]:
    """返回要导出的 (纪要, 逐字稿)；单独要求的那部分不存在时报 NothingToExport。"""
    minutes = fill_speakers(data.minutes_md, data.speakers).strip() if data.minutes_md else ""
    transcript = turns(data)
    want_minutes = content in ("minutes", "both")
    want_transcript = content in ("transcript", "both")
    if content == "minutes" and not minutes:
        raise NothingToExport("这场会议还没有纪要")
    if content == "transcript" and not transcript:
        raise NothingToExport("这场会议还没有逐字稿")
    if content == "both" and not minutes and not transcript:
        raise NothingToExport("这场会议还没有可以导出的内容")
    return (minutes if want_minutes and minutes else None), (transcript if want_transcript and transcript else None)


# ---------- Markdown 解析（纪要用到的常用子集）----------

FENCE = re.compile(r"^\s*(```|~~~)")
HEADING = re.compile(r"^(#{1,6})\s+(.*?)(?:\s+#+)?\s*$")
RULE = re.compile(r"^\s{0,3}([-*_])(?:\s*\1){2,}\s*$")
BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")
ORDERED = re.compile(r"^(\s*)(\d{1,3})[.)]\s+(.*)$")
QUOTE = re.compile(r"^\s{0,3}>\s?(.*)$")
TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(?:\|\s*:?-{2,}:?\s*)*\|?\s*$")
INLINE = re.compile(
    r"(?P<bold>\*\*(?=\S)(?P<bold_text>.+?)(?<=\S)\*\*)"
    r"|(?P<code>`(?P<code_text>[^`]+)`)"
    r"|(?P<italic>(?<!\*)\*(?=\S)(?P<italic_text>[^*\n]+?)(?<=\S)\*(?!\*))"
    r"|(?P<stamp>\[\d{1,2}:\d{2}(?::\d{2})?\])"
)
ESCAPED = re.compile(r"\\([\\`*_{}\[\]()#+\-.!|>~])")


@dataclass
class Block:
    kind: Literal["heading", "bullet", "ordered", "quote", "table", "para", "code"]
    text: str = ""
    level: int = 0  # 标题级别或列表缩进层级
    marker: str = ""
    rows: list[list[str]] = field(default_factory=list)


def _cells(line: str) -> list[str]:
    body = line.strip()
    body = body.removeprefix("|")
    if body.endswith("|") and not body.endswith("\\|"):
        body = body[:-1]
    return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", body)]


def _depth(indent: str) -> int:
    return min(len(indent.replace("\t", "    ")) // 2, 2)


def parse_markdown(text: str) -> list[Block]:
    blocks: list[Block] = []
    in_fence = False
    table: Block | None = None
    for line in text.splitlines():
        if FENCE.match(line):
            in_fence = not in_fence
            table = None
            continue
        if in_fence:
            blocks.append(Block("code", line))
            continue
        stripped = line.strip()
        if stripped.startswith("|"):
            if TABLE_SEP.match(stripped):
                continue
            if table is None:
                table = Block("table")
                blocks.append(table)
            table.rows.append(_cells(stripped))
            continue
        table = None
        if not stripped or RULE.match(line):
            continue
        if m := HEADING.match(line):
            blocks.append(Block("heading", m.group(2), level=len(m.group(1))))
        elif m := BULLET.match(line):
            blocks.append(Block("bullet", m.group(2), level=_depth(m.group(1))))
        elif m := ORDERED.match(line):
            blocks.append(Block("ordered", m.group(3), level=_depth(m.group(1)), marker=f"{m.group(2)}."))
        elif m := QUOTE.match(line):
            blocks.append(Block("quote", m.group(1).strip()))
        else:
            blocks.append(Block("para", stripped))
    return blocks


def inline_runs(text: str) -> Iterator[tuple[str, str]]:
    """把一行拆成 (样式, 文字)：plain / bold / italic / stamp。行内代码按普通文字处理。"""
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            yield "plain", ESCAPED.sub(r"\1", text[pos : m.start()])
        if m.group("bold"):
            yield "bold", ESCAPED.sub(r"\1", m.group("bold_text"))
        elif m.group("code"):
            yield "plain", m.group("code_text")
        elif m.group("italic"):
            yield "italic", ESCAPED.sub(r"\1", m.group("italic_text"))
        else:
            yield "stamp", m.group("stamp")
        pos = m.end()
    if pos < len(text):
        yield "plain", ESCAPED.sub(r"\1", text[pos:])


def plain_inline(text: str) -> str:
    return "".join(t for _, t in inline_runs(text))


def has_h1(markdown: str) -> bool:
    in_fence = False
    for line in markdown.splitlines():
        if FENCE.match(line):
            in_fence = not in_fence
        elif not in_fence and line.startswith("# "):
            return True
    return False


def demote_headings(markdown: str) -> str:
    """纪要里的标题整体降一级，好让文档标题（会议名）当一级标题。"""
    out: list[str] = []
    in_fence = False
    for line in markdown.splitlines():
        if FENCE.match(line):
            in_fence = not in_fence
        elif not in_fence and re.match(r"^#{1,5}\s", line):
            line = "#" + line
        out.append(line)
    return "\n".join(out)


# ---------- Markdown ----------


def _md_escape(text: str) -> str:
    return re.sub(r"([\\`*_<])", r"\\\1", text)


def render_markdown(data: ExportData, content: ExportContent) -> str:
    minutes, transcript = sections(data, content)
    parts = [f"# {data.title}", "\n".join(f"- {label}：{value}" for label, value in header_lines(data))]
    if minutes:
        if not has_h1(minutes):
            parts.append("## 会议纪要")
        parts.append(demote_headings(minutes))
    if transcript:
        parts.append("## 逐字稿")
        parts.extend(
            f"**{speaker_name(data.speakers, t.speaker)}** [{clock(t.start_ms)}]：{_md_escape(t.text)}"
            for t in transcript
        )
    return "\n\n".join(p for p in parts if p.strip()) + "\n"


# ---------- 纯文本 ----------


def markdown_to_text(markdown: str) -> str:
    lines: list[str] = []
    in_fence = False
    for line in markdown.splitlines():
        if FENCE.match(line):
            in_fence = not in_fence
            continue
        if in_fence:
            lines.append(line)
            continue
        stripped = line.strip()
        if stripped.startswith("|"):
            if not TABLE_SEP.match(stripped):
                lines.append(" | ".join(plain_inline(c) for c in _cells(stripped)))
        elif not stripped or RULE.match(line):
            lines.append("")
        elif m := HEADING.match(line):
            lines.append(plain_inline(m.group(2)))
        elif m := BULLET.match(line):
            lines.append("  " * _depth(m.group(1)) + "- " + plain_inline(m.group(2)))
        elif m := ORDERED.match(line):
            lines.append("  " * _depth(m.group(1)) + f"{m.group(2)}. " + plain_inline(m.group(3)))
        elif m := QUOTE.match(line):
            lines.append(plain_inline(m.group(1)))
        else:
            lines.append(plain_inline(stripped))
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def render_text(data: ExportData, content: ExportContent) -> str:
    minutes, transcript = sections(data, content)
    parts = [data.title, "\n".join(f"{label}：{value}" for label, value in header_lines(data))]
    if minutes:
        if not has_h1(minutes):
            parts.append("会议纪要")
        parts.append(markdown_to_text(minutes))
    if transcript:
        parts.append("逐字稿")
        parts.extend(f"{speaker_name(data.speakers, t.speaker)} [{clock(t.start_ms)}]：{t.text}" for t in transcript)
    return BOM + "\n\n".join(p for p in parts if p.strip()) + "\n"


# ---------- 字幕 ----------


def render_srt(data: ExportData) -> str:
    cues: list[str] = []
    for seg in data.segments:
        text = _clean(seg.text)
        if not text:
            continue
        end = max(seg.end_ms, seg.start_ms + 500)
        name = speaker_name(data.speakers, seg.speaker)
        cues.append(f"{len(cues) + 1}\n{srt_clock(seg.start_ms)} --> {srt_clock(end)}\n{name}：{text}\n")
    if not cues:
        raise NothingToExport("这场会议还没有逐字稿")
    return BOM + "\n".join(cues)


# ---------- Word ----------


def _use_fonts(rpr: Any, latin: str | None = None) -> None:
    fonts = rpr.get_or_add_rFonts()
    # 主题字体属性优先于显式字体，不删掉的话下面的设置不生效
    fonts.attrib.pop(qn("w:eastAsiaTheme"), None)
    fonts.set(qn("w:eastAsia"), CJK_FONT)
    if latin:
        for attr in ("ascii", "hAnsi"):
            fonts.attrib.pop(qn(f"w:{attr}Theme"), None)
            fonts.set(qn(f"w:{attr}"), latin)


def _setup_fonts(doc: Any) -> None:
    defaults = doc.styles.element.xpath("w:docDefaults/w:rPrDefault/w:rPr")
    if defaults:
        # 默认模板正文的西文是衬线的 Cambria，和中文黑体混排不协调，换成无衬线
        _use_fonts(defaults[0], LATIN_FONT)
        lang = defaults[0].find(qn("w:lang"))
        if lang is not None:
            # 引号等中西共用的标点按中文排版、用中文字体
            lang.set(qn("w:eastAsia"), "zh-CN")
    for style in doc.styles:
        rpr = style.element.rPr
        if rpr is not None and rpr.rFonts is not None and rpr.rFonts.get(qn("w:eastAsiaTheme")):
            _use_fonts(rpr)


def _add_runs(paragraph: DocxParagraph, text: str, bold: bool = False) -> None:
    for kind, value in inline_runs(text):
        if not value:
            continue
        run = paragraph.add_run(value)
        if bold or kind == "bold":
            run.bold = True
        if kind == "italic":
            run.italic = True
        if kind == "stamp":
            run.font.color.rgb = GRAY


def _add_table(doc: Any, rows: list[list[str]]) -> None:
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, row + [""] * (width - len(row)), strict=True):
            _add_runs(cell.paragraphs[0], value, bold=i == 0)
    doc.add_paragraph()


def _add_minutes(doc: Any, minutes: str) -> None:
    if not has_h1(minutes):
        doc.add_heading("会议纪要", level=1)
    for block in parse_markdown(minutes):
        if block.kind == "heading":
            _add_runs(doc.add_paragraph(style=f"Heading {block.level}"), block.text)
        elif block.kind == "bullet":
            style = "List Bullet" if block.level == 0 else f"List Bullet {block.level + 1}"
            _add_runs(doc.add_paragraph(style=style), block.text)
        elif block.kind == "ordered":
            # 不用 List Number：整篇共用一个编号，第二个列表会接着上一个往下数
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75 * (block.level + 1))
            p.paragraph_format.first_line_indent = Cm(-0.75)
            p.add_run(f"{block.marker} ")
            _add_runs(p, block.text)
        elif block.kind == "quote":
            _add_runs(doc.add_paragraph(style="Quote"), block.text)
        elif block.kind == "table":
            _add_table(doc, block.rows)
        elif block.kind == "code":
            doc.add_paragraph(block.text)
        else:
            _add_runs(doc.add_paragraph(), block.text)


def render_docx(data: ExportData, content: ExportContent) -> bytes:
    minutes, transcript = sections(data, content)
    doc = Document()
    _setup_fonts(doc)
    # 默认模板自带作者 python-docx 和 2013 年的时间，换成这份文档自己的
    props = doc.core_properties
    props.title, props.author, props.comments, props.language = data.title, "", "", "zh-CN"
    props.created = props.modified = datetime.now(UTC)
    doc.add_paragraph(data.title, style="Title")
    for label, value in header_lines(data):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Cm(0)
        p.add_run(f"{label}：").font.color.rgb = GRAY
        p.add_run(value)
    if minutes:
        _add_minutes(doc, minutes)
    if transcript:
        doc.add_heading("逐字稿", level=1)
        for t in transcript:
            p = doc.add_paragraph()
            p.add_run(speaker_name(data.speakers, t.speaker)).bold = True
            p.add_run(f" [{clock(t.start_ms)}]").font.color.rgb = GRAY
            p.add_run(f"：{t.text}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------- 文件 ----------


def render(data: ExportData, fmt: ExportFormat, content: ExportContent) -> bytes:
    if fmt == "docx":
        return render_docx(data, content)
    if fmt == "srt":
        return render_srt(data).encode()
    if fmt == "md":
        return render_markdown(data, content).encode()
    return render_text(data, content).encode()


_UNSAFE_NAME = re.compile(r'[\x00-\x1f\x7f\\/:*?"<>|]+')
SUFFIX = {"minutes": ("-纪要", "-minutes"), "transcript": ("-逐字稿", "-transcript"), "both": ("", "")}


def file_names(title: str, fmt: ExportFormat, content: ExportContent) -> tuple[str, str]:
    """返回 (文件名, 纯 ASCII 的备用文件名)。字幕只有逐字稿，不加后缀。"""
    stem = re.sub(r"\s+", " ", _UNSAFE_NAME.sub(" ", title)).strip(" .")[:80].strip(" .") or "会议记录"
    suffix, ascii_suffix = SUFFIX["both" if fmt == "srt" else content]
    name = f"{stem}{suffix}.{fmt}"
    fallback = name if name.isascii() else f"meeting{ascii_suffix}.{fmt}"
    return name, fallback


def content_disposition(name: str, fallback: str) -> str:
    # filename 给不认 RFC 5987 的老客户端，filename* 带上 UTF-8 编码的中文名
    return f"attachment; filename=\"{fallback}\"; filename*=UTF-8''{quote(name, safe='')}"
