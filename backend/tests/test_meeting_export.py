from __future__ import annotations

import io
import re
from datetime import UTC, datetime, timedelta, timezone
from urllib.parse import quote

import docx
import pytest
from docx.oxml.ns import qn

from app.db import utcnow
from app.meeting.export import ExportData, ExportSegment, file_names, render_markdown, resolve_speaker, turns
from app.models import Meeting, MeetingSegment
from app.security import new_job_id
from tests.conftest import add_user, login

SPEAKERS = {
    "S1": {"name": "张老师", "guess": None, "merged_into": None},
    "S2": {"name": "", "guess": {"name": "李同学", "evidence": "", "confidence": "high"}, "merged_into": None},
    "S3": {"name": "", "guess": None, "merged_into": "S1"},
}
MINUTES = """# 会议纪要：风场预测

## 概要
[[S1]] 主持，[[S2]] 汇报了 **baseline** 进展 [00:08]。

## 待办
| 事项 | 负责人 | 截止 | 时间戳 |
| --- | --- | --- | --- |
| 继续调参 | [[S2]] | 未明确 | [00:00:15] |

- [[S3]] 提供数据
- 未知的人 [[S9]]
1. 第一项
"""
# (开始, 结束, 说话人, 原文, 整理后)
SEGMENTS = [
    (0, 4000, "S1", "嗯大家好今天讨论风场预测", "大家好，今天讨论风场预测。"),
    (4000, 8000, "S3", "我补充一下数据情况", "我补充一下数据情况。"),
    (8000, 12000, "S2", "好的我这周跑完了", "好的，我这周跑完了 baseline"),
    (12000, 15000, "S2", "loss 降了", "loss 降了一些。"),
    (15000, 16000, "S1", "", ""),
    (3725000, 3730000, "S1", "那就这样", "那就这样。"),
]


def make_meeting(app, user_id: int, *, title="周三组会", minutes: str | None = MINUTES, segments=SEGMENTS) -> str:
    with app.state.ctx.Session() as db:
        m = Meeting(
            id=new_job_id(),
            user_id=user_id,
            title=title,
            filename="组会.m4a",
            status="done",
            duration_ms=3_730_000,
            provider_name="模拟识别",
            speakers=SPEAKERS,
            minutes_md=minutes,
            created_at=datetime(2026, 10, 4, 18, 0, tzinfo=UTC),  # 东八区已是 10 月 5 日
        )
        db.add(m)
        for i, (start, end, speaker, raw, text) in enumerate(segments):
            db.add(
                MeetingSegment(
                    meeting_id=m.id,
                    idx=i,
                    start_ms=start,
                    end_ms=end,
                    asr_speaker=speaker,
                    speaker=speaker,
                    raw_text=raw,
                    text=text,
                )
            )
        db.commit()
        return m.id


@pytest.fixture(autouse=True)
def east_eight(monkeypatch):
    # 日期按 BDW_TZ_OFFSET_HOURS 换算，固定成东八区，免得跑测试的机器设了别的时区
    monkeypatch.setattr("app.meeting.export.LOCAL_TZ", timezone(timedelta(hours=8)))


@pytest.fixture
def owner(app, client) -> int:
    user_id = add_user(app, "alice", "alice-password")
    login(client, "alice", "alice-password")
    return user_id


def export(client, mid: str, fmt: str, content: str = "both"):
    return client.get(f"/api/meetings/{mid}/export", params={"format": fmt, "content": content})


def disposition(resp) -> tuple[str, str]:
    value = resp.headers["content-disposition"]
    assert value.startswith("attachment;")
    fallback = re.search(r'filename="([^"]+)"', value).group(1)
    encoded = re.search(r"filename\*=UTF-8''(\S+)", value).group(1)
    return fallback, encoded


def test_markdown_both(app, client, owner):
    mid = make_meeting(app, owner)
    resp = export(client, mid, "md")
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == "text/markdown; charset=utf-8"
    assert disposition(resp) == ("meeting.md", quote("周三组会.md", safe=""))
    body = resp.text
    assert "[[S" not in body
    assert body.startswith("# 周三组会\n\n- 日期：2026-10-05\n- 时长：1 小时 2 分\n- 参会人：张老师、说话人 2\n")
    assert "- 识别服务：模拟识别" in body
    # 纪要标题降一级，挂在会议标题下面
    assert "\n## 会议纪要：风场预测\n" in body and "\n### 概要\n" in body and "\n### 待办\n" in body
    assert "张老师 主持，说话人 2 汇报了 **baseline** 进展 [00:08]。" in body
    assert "| 继续调参 | 说话人 2 | 未明确 | [00:00:15] |" in body
    assert "- 张老师 提供数据" in body and "未知的人 说话人 9" in body
    # 逐字稿：S3 合并到 S1，和 S1 相邻的话连成一段；用整理后的文字；英文之间补空格
    assert "\n## 逐字稿\n" in body
    assert "**张老师** [00:00:00]：大家好，今天讨论风场预测。我补充一下数据情况。" in body
    assert "**说话人 2** [00:00:08]：好的，我这周跑完了 baseline loss 降了一些。" in body
    assert "**张老师** [01:02:05]：那就这样。" in body
    assert "嗯大家好" not in body
    assert body.index("## 会议纪要") < body.index("## 逐字稿")


def test_markdown_minutes_and_transcript_only(app, client, owner):
    mid = make_meeting(app, owner)
    minutes = export(client, mid, "md", "minutes")
    assert minutes.status_code == 200
    assert "## 会议纪要：风场预测" in minutes.text and "逐字稿" not in minutes.text
    assert disposition(minutes) == ("meeting-minutes.md", quote("周三组会-纪要.md", safe=""))

    transcript = export(client, mid, "md", "transcript")
    assert transcript.status_code == 200
    assert "会议纪要" not in transcript.text and "## 逐字稿" in transcript.text
    assert disposition(transcript)[1] == quote("周三组会-逐字稿.md", safe="")


def test_plain_text(app, client, owner):
    mid = make_meeting(app, owner)
    resp = export(client, mid, "txt")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "text/plain; charset=utf-8"
    assert disposition(resp) == ("meeting.txt", quote("周三组会.txt", safe=""))
    assert resp.content.startswith("﻿".encode())
    body = resp.content.decode("utf-8-sig")
    assert body.startswith("周三组会\n\n日期：2026-10-05\n")
    assert "[[S" not in body and "**" not in body and "#" not in body and "---" not in body
    assert "\n会议纪要：风场预测\n" in body
    assert "张老师 主持，说话人 2 汇报了 baseline 进展 [00:08]。" in body
    assert "事项 | 负责人 | 截止 | 时间戳\n继续调参 | 说话人 2 | 未明确 | [00:00:15]" in body
    assert "- 张老师 提供数据" in body and "1. 第一项" in body
    assert "\n逐字稿\n" in body
    assert "张老师 [00:00:00]：大家好，今天讨论风场预测。我补充一下数据情况。" in body


def test_srt_ignores_content(app, client, owner):
    mid = make_meeting(app, owner)
    resp = export(client, mid, "srt")  # 前端默认 content=both，字幕照样只导逐字稿
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("application/x-subrip")
    assert disposition(resp) == ("meeting.srt", quote("周三组会.srt", safe=""))
    body = resp.content.decode("utf-8-sig")
    cues = body.strip().split("\n\n")
    assert len(cues) == 5  # 空文字那句跳过；字幕每句一条，不合并
    assert cues[0] == "1\n00:00:00,000 --> 00:00:04,000\n张老师：大家好，今天讨论风场预测。"
    assert cues[1] == "2\n00:00:04,000 --> 00:00:08,000\n张老师：我补充一下数据情况。"
    assert cues[2].endswith("说话人 2：好的，我这周跑完了 baseline")
    assert cues[4] == "5\n01:02:05,000 --> 01:02:10,000\n张老师：那就这样。"
    assert "会议纪要" not in body and "[[S" not in body
    assert export(client, mid, "srt", "minutes").status_code == 200


def all_docx_text(document) -> str:
    texts = [p.text for p in document.paragraphs]
    texts += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(texts)


def test_docx(app, client, owner):
    mid = make_meeting(app, owner)
    resp = export(client, mid, "docx")
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == ("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert disposition(resp) == ("meeting.docx", quote("周三组会.docx", safe=""))
    document = docx.Document(io.BytesIO(resp.content))
    text = all_docx_text(document)
    assert "[[S" not in text and "**" not in text and "| ---" not in text
    assert document.core_properties.title == "周三组会"

    paragraphs = document.paragraphs
    styles = {p.text: p.style.name for p in paragraphs}
    assert paragraphs[0].text == "周三组会" and paragraphs[0].style.name == "Title"
    assert paragraphs[1].text == "日期：2026-10-05"
    assert styles["会议纪要：风场预测"] == "Heading 1"
    assert styles["概要"] == "Heading 2"
    assert styles["逐字稿"] == "Heading 1"
    assert styles["张老师 提供数据"] == "List Bullet"
    assert "1. 第一项" in styles

    summary = next(p for p in paragraphs if p.text.startswith("张老师 主持"))
    assert summary.text == "张老师 主持，说话人 2 汇报了 baseline 进展 [00:08]。"
    assert any(r.bold and r.text == "baseline" for r in summary.runs)
    assert any(r.text == "[00:08]" and r.font.color.rgb is not None for r in summary.runs)

    [table] = document.tables
    rows = [[c.text for c in row.cells] for row in table.rows]
    assert rows == [["事项", "负责人", "截止", "时间戳"], ["继续调参", "说话人 2", "未明确", "[00:00:15]"]]
    assert all(r.bold for r in table.rows[0].cells[0].paragraphs[0].runs)

    turn = next(p for p in paragraphs if p.text.startswith("张老师 [00:00:00]"))
    assert turn.text == "张老师 [00:00:00]：大家好，今天讨论风场预测。我补充一下数据情况。"
    name, stamp = turn.runs[0], turn.runs[1]
    assert name.text == "张老师" and name.bold
    assert stamp.text == " [00:00:00]" and stamp.font.color.rgb is not None

    # 中文字体：默认字体和标题样式都换成了显式的中文字体，不再走主题字体
    defaults = document.styles.element.xpath("w:docDefaults/w:rPrDefault/w:rPr/w:rFonts")[0]
    assert defaults.get(qn("w:eastAsia")) == "微软雅黑" and defaults.get(qn("w:eastAsiaTheme")) is None
    heading_fonts = document.styles["Heading 1"].element.rPr.rFonts
    assert heading_fonts.get(qn("w:eastAsia")) == "微软雅黑"


def test_docx_minutes_only(app, client, owner):
    mid = make_meeting(app, owner)
    resp = export(client, mid, "docx", "minutes")
    assert resp.status_code == 200
    assert disposition(resp)[1] == quote("周三组会-纪要.docx", safe="")
    text = all_docx_text(docx.Document(io.BytesIO(resp.content)))
    assert "会议纪要：风场预测" in text and "逐字稿" not in text and "[00:00:00]" not in text


def test_missing_content(app, client, owner):
    no_minutes = make_meeting(app, owner, minutes=None)
    assert export(client, no_minutes, "md", "minutes").status_code == 409
    both = export(client, no_minutes, "md")
    assert both.status_code == 200 and "会议纪要" not in both.text and "## 逐字稿" in both.text

    empty = make_meeting(app, owner, minutes=None, segments=[])
    for fmt in ("docx", "md", "txt", "srt"):
        resp = export(client, empty, fmt)
        assert resp.status_code == 409 and "还没有" in resp.json()["detail"]

    minutes_only = make_meeting(app, owner, segments=[])
    assert export(client, minutes_only, "md", "transcript").status_code == 409
    assert export(client, minutes_only, "srt").status_code == 409
    resp = export(client, minutes_only, "docx")
    assert resp.status_code == 200
    # 没有逐字稿时，参会人按编号列出未被合并的说话人
    assert "参会人：张老师、说话人 2" in all_docx_text(docx.Document(io.BytesIO(resp.content)))


def test_access_and_validation(app, client, owner):
    mid = make_meeting(app, owner)
    assert export(client, mid, "pdf").status_code == 422
    assert export(client, mid, "md", "all").status_code == 422
    assert client.get(f"/api/meetings/{mid}/export").status_code == 422
    assert export(client, "nope", "md").status_code == 404

    with app.state.ctx.Session() as db:
        deleted = db.get(Meeting, make_meeting(app, owner))
        deleted.deleted_at = utcnow()
        db.commit()
        deleted_id = deleted.id
    assert export(client, deleted_id, "md").status_code == 404

    add_user(app, "bob", "bob-password")
    login(client, "bob", "bob-password")
    for fmt in ("docx", "md", "txt", "srt"):
        assert export(client, mid, fmt).status_code == 404

    client.cookies.clear()
    assert export(client, mid, "md").status_code == 401


def test_unsafe_title(app, client, owner):
    mid = make_meeting(app, owner, title='a/b:c*?"<>|  组会\n纪要')
    resp = export(client, mid, "md")
    assert resp.status_code == 200
    fallback, encoded = disposition(resp)
    assert fallback == "meeting.md"
    assert encoded == quote("a b c 组会 纪要.md", safe="")

    assert file_names("Weekly sync", "docx", "both") == ("Weekly sync.docx", "Weekly sync.docx")
    assert file_names(" ... ", "txt", "transcript") == ("会议记录-逐字稿.txt", "meeting-transcript.txt")
    assert file_names("组会", "srt", "minutes") == ("组会.srt", "meeting.srt")


def test_render_helpers():
    # 合并关系成环时不死循环
    looped = {"S1": {"name": "", "merged_into": "S2"}, "S2": {"name": "", "merged_into": "S1"}}
    assert resolve_speaker(looped, "S1") in {"S1", "S2"}

    data = ExportData(
        title="t",
        speakers={"S1": {"name": "Ann", "merged_into": None}},
        segments=[
            ExportSegment(0, 1000, "S1", "Hello"),
            ExportSegment(1000, 2000, "S1", "world.\n下一行 *星号* `代码` <b>"),
        ],
        minutes_md="## 只有二级标题\n\n```\n# 代码里的井号\n```\n",
    )
    [turn] = turns(data)
    assert turn.text == "Hello world. 下一行 *星号* `代码` <b>"
    md = render_markdown(data, "both")
    # 没有一级标题的纪要补一个“会议纪要”；代码块里的井号不降级；逐字稿里的 Markdown 符号转义
    assert "\n## 会议纪要\n\n### 只有二级标题\n" in md
    assert "\n# 代码里的井号\n" in md
    assert "**Ann** [00:00:00]：Hello world. 下一行 \\*星号\\* \\`代码\\` \\<b>" in md
